#!/usr/bin/env python3
"""
Bag labeler: step through every pick bag (final_approach_and_pick_*) recorded on Oct 6/7 2026 under /mnt/data and
~/apple_harvest_ws/data, watch both camera feeds and mark when the apple was picked or slipped.

  python3 label_bags.py [--out ../labels/PICK_TIMES_2026.xlsx] [--dates 2026-10-06 2026-10-07]

Buttons (keyboard shortcut in brackets)
  Mark pick [P] / Mark slip [S]  mark the current playback time; one event per bag, marking again replaces it
  Restart videos [R]             play both feeds from the start
  Next bag [N]                   save the mark and go to the next bag
  N/A [X]                        record the bag as N/A (left out of the spreadsheet) and go to the next bag
  Back [B]                       go to the previous bag to check or redo it
  Space plays/pauses, Left/Right step one frame, the slider scrubs.

Every decision is saved straight away to <out>.progress.json and the spreadsheet is rewritten from it, so you can
close the window at any time; the next run starts at the first bag that has no label yet. Times are seconds from the
bag's start (the same zero as plotting the bag's other topics), plus the absolute epoch time.
"""

import argparse
import bisect
import datetime
import json
import os
import re
import sys
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, asdict
from pathlib import Path

import cv2
import numpy as np
import yaml
from PyQt5 import QtCore, QtGui, QtWidgets

import rosbag2_py
from rclpy.serialization import deserialize_message
from sensor_msgs.msg import Image

# pip's opencv-python points Qt at its own bundled plugins, which don't work with the system PyQt5
# ("Could not load the Qt platform plugin xcb"); use PyQt5's instead
for _var in ('QT_QPA_PLATFORM_PLUGIN_PATH', 'QT_QPA_FONTDIR'):
    if 'cv2' in os.environ.get(_var, ''):
        del os.environ[_var]

DEFAULT_ROOTS = ['/mnt/data', str(Path.home() / 'apple_harvest_ws' / 'data')]
PICK_ANALYSIS = Path(__file__).resolve().parent.parent  # pick_analysis/: scripts/, labels/, plots/, videos/
DEFAULT_OUT = str(PICK_ANALYSIS / 'labels' / 'PICK_TIMES_2026.xlsx')
BAG_PREFIX = 'final_approach_and_pick'
MAX_FRAME_WIDTH = 640  # frames are downscaled to this when loaded, to keep a few bags in memory
SPEEDS = [0.1, 0.25, 0.5, 1.0, 2.0]

# encoding -> (channels, cv2 conversion to RGB or None)
CONVERSIONS = {
    'rgb8': (3, None),
    'bgr8': (3, cv2.COLOR_BGR2RGB),
    'rgba8': (4, cv2.COLOR_RGBA2RGB),
    'bgra8': (4, cv2.COLOR_BGRA2RGB),
    'mono8': (1, cv2.COLOR_GRAY2RGB),
    'yuv422_yuy2': (2, cv2.COLOR_YUV2RGB_YUY2),
    'yuyv': (2, cv2.COLOR_YUV2RGB_YUY2),
    'yuv422': (2, cv2.COLOR_YUV2RGB_UYVY),
    'uyvy': (2, cv2.COLOR_YUV2RGB_UYVY),
}

COLUMNS = ['Bag filename', 'Date captured', 'Batch', 'Apple', 'Pull type', 'Pick/Slip',
           'Pick/Slip time (s from bag start)', 'Pick/Slip time (epoch s)', 'Bag path']


@dataclass
class BagInfo:
    path: str
    name: str
    storage_id: str
    start_ns: int
    duration_s: float
    batch: int | None
    apple: int | None
    pull_type: str
    pick_result: str
    image_topics: list

    @property
    def start(self):
        return datetime.datetime.fromtimestamp(self.start_ns / 1e9)


def _load_yaml(path):
    try:
        with open(path) as f:
            return yaml.safe_load(f) or {}
    except (OSError, yaml.YAMLError):
        return {}


def read_bag_info(bag_dir):
    meta = _load_yaml(bag_dir / 'metadata.yaml').get('rosbag2_bagfile_information')
    if not meta:
        return None
    topics = [t['topic_metadata']['name'] for t in meta.get('topics_with_message_count', [])
              if t['topic_metadata']['type'] == 'sensor_msgs/msg/Image' and t.get('message_count', 0) > 0]
    # RealSense first, then the gripper camera
    topics.sort(key=lambda t: ('realsense' not in t, t))

    pick_meta = _load_yaml(bag_dir.parent / 'pick_metadata.yaml')
    pull_type = pick_meta.get('pull_pattern', '')
    if not pull_type:
        for batch_meta in bag_dir.parent.parent.glob('batch_*_metadata.yaml'):
            pull_type = _load_yaml(batch_meta).get('pull_pattern', '')

    batch = re.search(r'batch_(\d+)', str(bag_dir))
    apple = re.search(r'apple_(\d+)', str(bag_dir))
    return BagInfo(
        path=str(bag_dir),
        name=bag_dir.name,
        storage_id=meta.get('storage_identifier', 'sqlite3'),
        start_ns=meta['starting_time']['nanoseconds_since_epoch'],
        duration_s=meta['duration']['nanoseconds'] / 1e9,
        batch=int(batch.group(1)) if batch else None,
        apple=int(apple.group(1)) if apple else None,
        pull_type=pull_type,
        pick_result=str(pick_meta.get('pick_result', '')),
        image_topics=topics,
    )


def find_bags(roots, dates):
    bags = {}
    for root in roots:
        for dirpath, dirnames, filenames in os.walk(root, onerror=lambda e: None):
            dirnames[:] = sorted(d for d in dirnames if not d.startswith('.') and d != 'lost+found')
            if 'metadata.yaml' not in filenames:
                continue
            dirnames[:] = []  # a bag directory only holds its .db3 files
            name = os.path.basename(dirpath)
            if not name.startswith(BAG_PREFIX) or name in bags:
                continue
            info = read_bag_info(Path(dirpath))
            if info and info.start.date() in dates:
                bags[name] = info
    return sorted(bags.values(), key=lambda b: b.start_ns)


def to_rgb(msg, max_width):
    enc = msg.encoding.lower()
    if enc not in CONVERSIONS:
        return None
    channels, conversion = CONVERSIONS[enc]
    h, w = msg.height, msg.width
    rows = np.frombuffer(msg.data, np.uint8).reshape(h, msg.step)[:, :w * channels]
    img = rows.reshape(h, w, channels) if channels > 1 else rows
    if conversion is not None:
        img = cv2.cvtColor(img, conversion)
    if w > max_width:
        img = cv2.resize(img, (max_width, round(h * max_width / w)), interpolation=cv2.INTER_AREA)
    return img if img.base is None else img.copy()  # don't keep the whole message buffer alive


def load_frames(bag):
    """Returns {topic: (np.array of receive times in ns, [RGB frames])}."""
    reader = rosbag2_py.SequentialReader()
    reader.open(rosbag2_py.StorageOptions(uri=bag.path, storage_id=bag.storage_id),
                rosbag2_py.ConverterOptions('', ''))
    reader.set_filter(rosbag2_py.StorageFilter(topics=bag.image_topics))
    frames = {t: ([], []) for t in bag.image_topics}
    while reader.has_next():
        topic, data, t_ns = reader.read_next()
        img = to_rgb(deserialize_message(data, Image), MAX_FRAME_WIDTH)
        if img is not None:
            frames[topic][0].append(t_ns)
            frames[topic][1].append(img)
    return {t: (np.array(ts, dtype=np.int64), imgs) for t, (ts, imgs) in frames.items() if ts}


class LabelStore:
    """Labels keyed by bag path, saved to JSON; the spreadsheet is regenerated from it after every change."""

    def __init__(self, xlsx_path):
        self.xlsx_path = Path(xlsx_path)
        self.json_path = self.xlsx_path.with_suffix('.progress.json')
        self.labels = json.loads(self.json_path.read_text()) if self.json_path.exists() else {}

    def get(self, bag):
        return self.labels.get(bag.path)

    def set(self, bag, label, time_ns=None):
        entry = {'label': label, 'labeled_at': datetime.datetime.now().isoformat(timespec='seconds'),
                 'bag': asdict(bag)}
        if time_ns is not None:
            entry['time_ns'] = int(time_ns)
            entry['time_s'] = (time_ns - bag.start_ns) / 1e9
        self.labels[bag.path] = entry
        self.save()

    def save(self):
        tmp = self.json_path.with_suffix('.tmp')
        tmp.write_text(json.dumps(self.labels, indent=1))
        os.replace(tmp, self.json_path)
        self.write_xlsx()

    def write_xlsx(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = self.xlsx_path.stem[:31]  # the tab is named after the file, e.g. PICK_TIMES_2026
        ws.append(COLUMNS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        entries = sorted((e for e in self.labels.values() if e['label'] != 'na'),
                         key=lambda e: e['bag']['start_ns'])
        for e in entries:
            b = e['bag']
            ws.append([b['name'], datetime.datetime.fromtimestamp(b['start_ns'] / 1e9), b['batch'], b['apple'],
                       b['pull_type'], e['label'], round(e['time_s'], 3), round(e['time_ns'] / 1e9, 3), b['path']])
            ws.cell(ws.max_row, 2).number_format = 'yyyy-mm-dd hh:mm:ss'
            ws.cell(ws.max_row, 8).number_format = '0.000'
        for col, width in zip('ABCDEFGHI', [44, 20, 7, 7, 24, 10, 16, 18, 100]):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        tmp = self.xlsx_path.with_name('.' + self.xlsx_path.name + '.tmp')
        wb.save(tmp)
        os.replace(tmp, self.xlsx_path)


class VideoView(QtWidgets.QLabel):
    def __init__(self):
        super().__init__()
        self.setAlignment(QtCore.Qt.AlignCenter)
        self.setSizePolicy(QtWidgets.QSizePolicy.Ignored, QtWidgets.QSizePolicy.Ignored)
        self.setMinimumSize(320, 240)
        self.setStyleSheet('background: black; color: #aaa;')
        self._img = None

    def show_frame(self, img):
        self._img = img
        self._redraw()

    def show_text(self, text):
        self._img = None
        self.clear()
        self.setText(text)

    def _redraw(self):
        if self._img is None:
            return
        h, w = self._img.shape[:2]
        qimg = QtGui.QImage(self._img.data, w, h, 3 * w, QtGui.QImage.Format_RGB888)
        self.setPixmap(QtGui.QPixmap.fromImage(qimg).scaled(
            self.size(), QtCore.Qt.KeepAspectRatio, QtCore.Qt.SmoothTransformation))

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._redraw()


class Labeler(QtWidgets.QWidget):
    def __init__(self, bags, store):
        super().__init__()
        self.bags = bags
        self.store = store
        self.pool = ThreadPoolExecutor(max_workers=1)
        self.cache = {}  # bag index -> Future of load_frames
        self.idx = None
        self.frames = None
        self.mark = None  # (label, time_ns) for the current bag
        self.playing = False
        self.pos_ns = 0
        self.wall = QtCore.QElapsedTimer()

        self.setWindowTitle('Bag labeler')
        self.resize(1400, 750)
        self._build_ui()

        self.play_timer = QtCore.QTimer(self, interval=30, timeout=self._tick)
        self.load_timer = QtCore.QTimer(self, interval=100, timeout=self._check_loaded)

        first = next((i for i, b in enumerate(bags) if store.get(b) is None), len(bags) - 1)
        self.go_to(first)

    def _build_ui(self):
        self.info = QtWidgets.QLabel()
        self.info.setTextInteractionFlags(QtCore.Qt.TextSelectableByMouse)
        self.progress = QtWidgets.QLabel()
        self.progress.setAlignment(QtCore.Qt.AlignRight | QtCore.Qt.AlignVCenter)
        top = QtWidgets.QHBoxLayout()
        top.addWidget(self.info, 1)
        top.addWidget(self.progress)

        self.views = [VideoView(), VideoView()]
        self.view_titles = [QtWidgets.QLabel(), QtWidgets.QLabel()]
        videos = QtWidgets.QHBoxLayout()
        for title, view in zip(self.view_titles, self.views):
            col = QtWidgets.QVBoxLayout()
            col.addWidget(title)
            col.addWidget(view, 1)
            videos.addLayout(col, 1)

        self.play_btn = QtWidgets.QPushButton('Pause [Space]', clicked=self.toggle_play)
        self.slider = QtWidgets.QSlider(QtCore.Qt.Horizontal)
        self.slider.sliderMoved.connect(self._seek_ms)
        self.slider.sliderPressed.connect(lambda: self._seek_ms(self.slider.value()))
        self.time_label = QtWidgets.QLabel()
        self.time_label.setMinimumWidth(190)
        self.speed = QtWidgets.QComboBox()
        self.speed.addItems([f'{s:g}x' for s in SPEEDS])
        self.speed.setCurrentIndex(SPEEDS.index(1.0))
        self.speed.currentIndexChanged.connect(lambda _: self.wall.restart())
        transport = QtWidgets.QHBoxLayout()
        transport.addWidget(self.play_btn)
        transport.addWidget(self.slider, 1)
        transport.addWidget(self.time_label)
        transport.addWidget(QtWidgets.QLabel('Speed'))
        transport.addWidget(self.speed)

        self.mark_label = QtWidgets.QLabel()
        font = self.mark_label.font()
        font.setPointSize(font.pointSize() + 2)
        font.setBold(True)
        self.mark_label.setFont(font)
        self.mark_label.setMinimumWidth(420)
        self.status = QtWidgets.QLabel()
        self.status.setStyleSheet('color: #c62828;')

        buttons = QtWidgets.QHBoxLayout()
        self.action_buttons = []
        for text, slot, color in [('Back [B]', self.back, None),
                                  ('Mark pick [P]', lambda: self.set_mark('pick'), '#2e7d32'),
                                  ('Mark slip [S]', lambda: self.set_mark('slip'), '#ef6c00'),
                                  ('Restart videos [R]', self.restart, None),
                                  ('Next bag [N]', self.next_bag, '#1565c0'),
                                  ('N/A [X]', self.mark_na, '#616161')]:
            btn = QtWidgets.QPushButton(text, clicked=slot)
            btn.setMinimumHeight(44)
            if color:
                btn.setStyleSheet(f'QPushButton {{ background: {color}; color: white; font-weight: bold; }}')
            buttons.addWidget(btn)
            self.action_buttons.append(btn)

        # nothing takes keyboard focus, so Space and the arrows always go to the shortcuts below
        for w in self.action_buttons + [self.play_btn, self.slider, self.speed]:
            w.setFocusPolicy(QtCore.Qt.NoFocus)
        for key, slot in [('Space', self.toggle_play), ('Left', lambda: self.step(-1)),
                          ('Right', lambda: self.step(1)), ('P', lambda: self.set_mark('pick')),
                          ('S', lambda: self.set_mark('slip')), ('R', self.restart), ('N', self.next_bag),
                          ('X', self.mark_na), ('B', self.back)]:
            QtWidgets.QShortcut(QtGui.QKeySequence(key), self, slot)

        layout = QtWidgets.QVBoxLayout(self)
        layout.addLayout(top)
        layout.addLayout(videos, 1)
        layout.addLayout(transport)
        mark_row = QtWidgets.QHBoxLayout()
        mark_row.addWidget(self.mark_label)
        mark_row.addStretch(1)
        mark_row.addWidget(self.status)
        layout.addLayout(mark_row)
        layout.addLayout(buttons)

    # ---------- bag navigation ----------

    def _load(self, i):
        if 0 <= i < len(self.bags) and i not in self.cache:
            self.cache[i] = self.pool.submit(load_frames, self.bags[i])

    def go_to(self, i):
        self.pause()
        self.idx = i
        self.frames = None
        bag = self.bags[i]
        saved = self.store.get(bag)
        self.mark = (saved['label'], saved['time_ns']) if saved and saved['label'] != 'na' else None

        self._load(i)
        self._load(i + 1)  # prefetch while this one is being labeled
        for j in list(self.cache):
            if j not in (i - 1, i, i + 1):
                self.cache.pop(j).cancel()

        batch = bag.batch if bag.batch is not None else '?'
        apple = bag.apple if bag.apple is not None else '?'
        self.info.setText(
            f'<b>{bag.name}</b> &nbsp; batch {batch}, apple {apple} &nbsp;|&nbsp; pull type: <b>{bag.pull_type or "?"}'
            f'</b> &nbsp;|&nbsp; recorded result: {bag.pick_result or "?"} &nbsp;|&nbsp; '
            f'{bag.start:%Y-%m-%d %H:%M:%S}, {bag.duration_s:.1f} s<br><small>{bag.path}</small>')
        self._update_progress()
        self._update_mark_label(saved)
        self.status.clear()
        for title, view, topic in zip(self.view_titles, self.views, bag.image_topics + [None, None]):
            title.setText(topic or '(no image topic)')
            view.show_text('Loading…' if topic else '')
        self.load_timer.start()

    def _check_loaded(self):
        future = self.cache.get(self.idx)
        if future is None or not future.done():
            return
        self.load_timer.stop()
        try:
            self.frames = future.result()
        except Exception as e:  # unreadable bag: show why, it can still be marked N/A
            for view in self.views:
                view.show_text(f'Could not read bag:\n{e}')
            return
        if not self.frames:
            for view in self.views:
                view.show_text('No decodable image messages in this bag')
            return
        self.all_ts = np.unique(np.concatenate([ts for ts, _ in self.frames.values()]))
        self.t_first, self.t_last = int(self.all_ts[0]), int(self.all_ts[-1])
        self.slider.setRange(0, (self.t_last - self.t_first) // 1_000_000)
        self.restart()

    def _update_progress(self):
        labels = [self.store.get(b) for b in self.bags]
        done = sum(1 for l in labels if l)
        na = sum(1 for l in labels if l and l['label'] == 'na')
        self.progress.setText(f'Bag <b>{self.idx + 1}</b> of {len(self.bags)}<br>'
                              f'{done} labeled ({done - na} pick/slip, {na} N/A)')

    def _update_mark_label(self, saved=None):
        if self.mark:
            label, t_ns = self.mark
            text = f'Marked: {label.upper()} at {(t_ns - self.bags[self.idx].start_ns) / 1e9:.3f} s'
        elif saved and saved['label'] == 'na':
            text = 'Saved as N/A'
        else:
            text = 'Not marked'
        self.mark_label.setText(text)

    def set_mark(self, label):
        if not self.frames:
            return
        self.pause()
        self.mark = (label, self.pos_ns)
        self.status.clear()
        self._update_mark_label()

    def next_bag(self):
        if self.mark is None:
            saved = self.store.get(self.bags[self.idx])
            if not (saved and saved['label'] == 'na'):
                self.status.setText('Mark pick or slip first, or press N/A')
                return
        else:
            self.store.set(self.bags[self.idx], *self.mark)
        self._advance()

    def mark_na(self):
        self.store.set(self.bags[self.idx], 'na')
        self._advance()

    def _advance(self):
        if self.idx + 1 < len(self.bags):
            self.go_to(self.idx + 1)
        else:
            self._update_progress()
            self._update_mark_label(self.store.get(self.bags[self.idx]))
            remaining = sum(1 for b in self.bags if self.store.get(b) is None)
            self.status.setText('Last bag done.' + (f' {remaining} earlier bags still unlabeled.' if remaining else ''))

    def back(self):
        if self.idx > 0:
            self.go_to(self.idx - 1)

    # ---------- playback ----------

    def restart(self):
        if not self.frames:
            return
        self._set_pos(self.t_first)
        self.play()

    def play(self):
        if not self.frames:
            return
        if self.pos_ns >= self.t_last:
            self._set_pos(self.t_first)
        self.playing = True
        self.play_btn.setText('Pause [Space]')
        self.wall.restart()
        self.play_timer.start()

    def pause(self):
        self.playing = False
        self.play_btn.setText('Play [Space]')
        self.play_timer.stop()

    def toggle_play(self):
        self.pause() if self.playing else self.play()

    def step(self, direction):
        if not self.frames:
            return
        self.pause()
        if direction > 0:
            i = bisect.bisect_right(self.all_ts, self.pos_ns)
        else:
            i = bisect.bisect_left(self.all_ts, self.pos_ns) - 1
        self._set_pos(int(self.all_ts[min(max(i, 0), len(self.all_ts) - 1)]))

    def _seek_ms(self, ms):
        if self.frames:
            self._set_pos(self.t_first + ms * 1_000_000)
            self.wall.restart()

    def _tick(self):
        elapsed_ns = self.wall.nsecsElapsed()
        self.wall.restart()
        pos = self.pos_ns + int(elapsed_ns * SPEEDS[self.speed.currentIndex()])
        if pos >= self.t_last:
            pos = self.t_last
            self.pause()
        self._set_pos(pos)

    def _set_pos(self, pos_ns):
        self.pos_ns = min(max(pos_ns, self.t_first), self.t_last)
        for view, topic in zip(self.views, self.bags[self.idx].image_topics):
            if topic not in self.frames:
                continue
            ts, imgs = self.frames[topic]
            view.show_frame(imgs[max(bisect.bisect_right(ts, self.pos_ns) - 1, 0)])
        self.slider.blockSignals(True)
        self.slider.setValue((self.pos_ns - self.t_first) // 1_000_000)
        self.slider.blockSignals(False)
        t = (self.pos_ns - self.bags[self.idx].start_ns) / 1e9
        end = (self.t_last - self.bags[self.idx].start_ns) / 1e9
        self.time_label.setText(f't = {t:.3f} s / {end:.2f} s')

    def closeEvent(self, event):
        self.pool.shutdown(wait=False, cancel_futures=True)
        super().closeEvent(event)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--roots', nargs='+', default=DEFAULT_ROOTS, help='directories searched for bags')
    parser.add_argument('--dates', nargs='+', default=['2026-10-06', '2026-10-07'],
                        help='only bags started on these (local) dates')
    parser.add_argument('--out', default=DEFAULT_OUT, help='spreadsheet to write (progress JSON goes next to it)')
    parser.add_argument('--list', action='store_true', help='print the bags that would be labeled and exit')
    args = parser.parse_args()

    dates = {datetime.date.fromisoformat(d) for d in args.dates}
    print('Searching for bags in', ', '.join(args.roots), '...', flush=True)
    bags = find_bags(args.roots, dates)
    print(f'Found {len(bags)} pick bags from {", ".join(sorted(args.dates))}')
    if args.list:
        for b in bags:
            print(f'{b.start:%Y-%m-%d %H:%M:%S}  batch {b.batch}  apple {b.apple}  {b.pull_type:24s} {b.path}')
        return
    if not bags:
        sys.exit('Nothing to label')

    app = QtWidgets.QApplication(sys.argv)
    window = Labeler(bags, LabelStore(args.out))
    window.show()
    code = app.exec_()
    os._exit(code)  # labels are already saved; don't wait for a prefetch still reading a bag


if __name__ == '__main__':
    main()
